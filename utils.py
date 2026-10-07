"""Чтение XLSX, индекс строк и гибридный поиск. Python 3.11–3.12."""
import functools
import hashlib
import math
import re
import threading
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, time
from io import BytesIO
from urllib.parse import urlsplit

import numpy as np
import pandas as pd
import streamlit as st

MODEL_ID = "skatzR/USER-BGE-M3-MiniLM-L12-v2-Distilled"
MAX_XLSX_BYTES = 10 * 1024 * 1024
MAX_XLSX_UNPACKED_BYTES = 50 * 1024 * 1024
MAX_XLSX_ROWS = 20000
MAX_XLSX_COLUMNS = 100
MAX_XLSX_SHEETS = 20
MAX_XLSX_CELLS = 300000
TITLE_ALIASES = {"фраза", "запрос", "обращение", "вопрос", "заголовок", "phrase", "query", "title"}
DATE_ALIASES = {"дата", "date", "дата обращения", "дата создания"}
TOKEN_RE = re.compile(r"\w+", re.UNICODE)


@st.cache_resource(show_spinner=False)
def get_model(model_id=MODEL_ID):
    from sentence_transformers import SentenceTransformer
    return SentenceTransformer(model_id, trust_remote_code=False)


@st.cache_resource(show_spinner=False)
def model_lock():
    # Модель и токенизатор используются несколькими сессиями Streamlit.
    return threading.RLock()


@st.cache_resource(show_spinner=False)
def get_morph():
    import pymorphy3
    return pymorphy3.MorphAnalyzer()


def normalize_spaces(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def preprocess(text):
    return normalize_spaces(text).casefold().replace("ё", "е")



def tokenize(text):
    return tuple(TOKEN_RE.findall(preprocess(text)))


@functools.lru_cache(maxsize=50000)
def lemmatize_cached(word):
    return get_morph().parse(word)[0].normal_form if word.isalpha() else word


def lemmatize(word):
    return lemmatize_cached(word)


def safe_http_url(value):
    """Проверяет адрес для отображения ссылки; ничего не скачивает."""
    value = str(value or "").strip(" ")
    if not value or re.search(r"""[\s\x00-\x1f\x7f<>"'\\]""", value):
        return None
    try:
        parts = urlsplit(value)
        if (parts.scheme.lower() not in {"http", "https"} or not parts.hostname
                or parts.username is not None or parts.password is not None):
            return None
        _ = parts.port  # отклоняет некорректный порт
        return value
    except ValueError:
        return None


def excel_value(value):
    """Сохраняет нули и переносы строк, приводит даты к читаемому виду."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y" if value.time() == time() else "%d.%m.%Y %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, bool):
        return "Да" if value else "Нет"
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer():
            return str(int(value))
    return str(value).strip()


def excel_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def suggest_header_row(rows):
    """Предлагает строку заголовков; её всегда можно изменить вручную."""
    best = (float("-inf"), 1)
    aliases = TITLE_ALIASES | DATE_ALIASES | {"интент", "интенты", "комментарий", "comment", "intent"}
    for number, row in enumerate(rows[:50], 1):
        values = [normalize_spaces(excel_value(v)) for v in row if excel_value(v)]
        if not values:
            continue
        known = sum(preprocess(v) in aliases for v in values)
        short = sum(len(v) <= 60 for v in values)
        unique = len(set(preprocess(v) for v in values))
        score = known * 10 + min(len(values), 10) + short / len(values) + unique / len(values)
        if score > best[0]:
            best = (score, number)
    return best[1]


@st.cache_data(show_spinner=False, max_entries=8)
def read_xlsx(content):
    """Читает только XLSX. Формулы не вычисляет: берёт сохранённые значения."""
    from openpyxl import load_workbook
    if not content or len(content) > MAX_XLSX_BYTES:
        raise ValueError("Пустой XLSX или размер больше 10 МБ.")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > MAX_XLSX_UNPACKED_BYTES:
                raise ValueError("Книга слишком велика после распаковки (лимит 50 МБ).")
        formulas = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Не удалось открыть XLSX. Нужна обычная книга без пароля.") from exc
    values = None
    try:
        values = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        if len(formulas.worksheets) > MAX_XLSX_SHEETS:
            raise ValueError(f"В книге больше {MAX_XLSX_SHEETS} листов.")
        sheets, budget = {}, 0
        for ws in formulas.worksheets:
            if (ws.max_row or 0) > MAX_XLSX_ROWS or (ws.max_column or 0) > MAX_XLSX_COLUMNS:
                raise ValueError(f"Лист «{ws.title}»: лимит {MAX_XLSX_ROWS} строк и {MAX_XLSX_COLUMNS} колонок. Удалите лишние пустые строки/колонки и их форматирование.")
            rows, missing, errors = [], 0, 0
            for formula_row, cached_row in zip(ws.iter_rows(), values[ws.title].iter_rows()):
                if len(rows) >= MAX_XLSX_ROWS or len(formula_row) > MAX_XLSX_COLUMNS:
                    raise ValueError(f"Лист «{ws.title}» превышает лимит строк или колонок.")
                budget += len(formula_row)
                if budget > MAX_XLSX_CELLS:
                    raise ValueError(f"В книге больше {MAX_XLSX_CELLS} ячеек. Разделите файл.")
                row = []
                for cell, cached in zip(formula_row, cached_row):
                    value = cached.value
                    if cell.data_type == "f" and value is None:
                        missing += 1
                    if cached.data_type == "e":
                        errors += 1
                        value = None
                    row.append(value)
                rows.append(tuple(row))
            while rows and not any(excel_value(v) for v in rows[-1]):
                rows.pop()
            sheets[ws.title] = {"rows": tuple(rows), "header_row": suggest_header_row(rows),
                               "missing_formulas": missing, "error_cells": errors}
        return sheets
    finally:
        formulas.close()
        if values is not None:
            values.close()


@st.cache_data(show_spinner=False, max_entries=32)
def xlsx_table(sheet, header_row=1):
    rows = sheet["rows"]
    if not isinstance(header_row, int) or not 1 <= header_row <= max(1, len(rows)):
        raise ValueError("Строка заголовков отсутствует на листе.")
    if not rows:
        return {"columns": [], "records": [], "date_columns": []}
    width = max((i + 1 for row in rows[header_row - 1:] for i, v in enumerate(row) if excel_value(v)), default=0)
    header = rows[header_row - 1]
    columns, used = [], set()
    for i in range(width):
        original = normalize_spaces(excel_value(header[i])) if i < len(header) else ""
        base = original or f"Колонка {i + 1}"
        label, suffix = base, 2
        while label in used:
            label = f"{base} ({suffix})"
            suffix += 1
        used.add(label)
        columns.append({"id": f"c{i}", "label": label, "original": original})
    records, date_cols = [], set()
    for row_number, row in enumerate(rows[header_row:], header_row + 1):
        cells, dates = {}, {}
        for i, column in enumerate(columns):
            value = row[i] if i < len(row) else None
            cells[column["id"]] = excel_value(value)
            parsed_date = excel_date(value)
            if parsed_date is not None:
                dates[column["id"]] = parsed_date
                date_cols.add(column["id"])
        if any(cells.values()):
            records.append({"row_number": row_number, "cells": cells, "dates": dates})
    return {"columns": columns, "records": records,
            "date_columns": [c["id"] for c in columns if c["id"] in date_cols]}


def suggest_xlsx_columns(table):
    columns = table["columns"]
    title = next((c["id"] for c in columns if preprocess(c["original"]) in TITLE_ALIASES), None)
    if title is None:
        title = next((c["id"] for c in columns if c["id"] not in table["date_columns"]
                      and preprocess(c["original"]) not in DATE_ALIASES), None)
    title = title or (columns[0]["id"] if columns else None)
    search = [c["id"] for c in columns if c["id"] not in table["date_columns"]
              and preprocess(c["original"]) not in DATE_ALIASES]
    return {"title_column": title, "search_columns": search or ([title] if title else [])}


def parse_xlsx_cases(cfg):
    if cfg.get("kind") != "xlsx" or "table" not in cfg or "sheet_name" not in cfg:
        raise ValueError("Источник должен быть настроенным листом XLSX.")
    table = cfg["table"]
    columns = {c["id"]: c["label"] for c in table["columns"]}
    display_labels = {c["id"]: c.get("display_label", c["label"]) for c in table["columns"]}
    title_col = cfg.get("title_column")
    search_cols = list(cfg.get("search_columns", list(columns)))
    display_cols = list(cfg.get("display_columns", list(columns)))
    if (title_col is not None and title_col not in columns) or any(c not in columns for c in search_cols + display_cols):
        raise ValueError("Выбрана отсутствующая колонка.")
    if not search_cols:
        raise ValueError("Выберите хотя бы одну колонку для поиска.")
    date_col, start, end = cfg.get("date_column"), cfg.get("date_from"), cfg.get("date_to")
    if date_col is not None and date_col not in columns:
        raise ValueError("Колонка даты отсутствует.")
    if start and end and start > end:
        raise ValueError("Начальная дата позже конечной.")
    sheet_id = hashlib.sha256(cfg["sheet_name"].encode("utf-8")).hexdigest()[:16]
    cases = []
    for record in table["records"]:
        cells, row_number = record["cells"], record["row_number"]
        if date_col is not None:
            value = record["dates"].get(date_col)
            if value is None or (start and value < start) or (end and value > end):
                continue
        if not any(cells[c] for c in search_cols):
            continue
        title = (cells.get(title_col, "") if title_col else "") or f"Строка {row_number}"
        text = "\n".join(f"{columns[c]}: {cells[c]}" for c in search_cols if cells[c])
        uid = f'{cfg.get("source_id", cfg["name"])}::sheet:{sheet_id}::row:{row_number}'
        cases.append({"case_uid": uid, "source_file": cfg["name"], "sheet_name": cfg["sheet_name"],
                      "row_number": row_number, "title": title, "search_text": text,
                      "search_proc": preprocess(text), "fields": [
                          {"label": display_labels[c], "value": cells[c]} for c in display_cols
                          if cells[c] and c != title_col]})
    return cases


@st.cache_data(show_spinner=False, max_entries=32)
def build_database(workbook_configs):
    """Индексирует только строки настроенных листов XLSX."""
    rows, identities = [], set()
    for cfg in workbook_configs:
        if cfg.get("kind") != "xlsx" or "table" not in cfg or "sheet_name" not in cfg:
            raise ValueError("Для индекса нужны настроенные листы XLSX.")
        identity = (cfg.get("source_id", cfg["name"]), cfg["sheet_name"])
        if identity in identities:
            raise ValueError("Один лист одного источника выбран несколько раз.")
        identities.add(identity)
        rows.extend(parse_xlsx_cases(cfg))
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    texts = tuple(row["search_text"] for row in rows)
    tokens = tuple(tokenize(text) for text in texts)
    lemmas = tuple(tuple(map(lemmatize_cached, words)) for words in tokens)
    postings, lemma_postings = defaultdict(set), defaultdict(set)
    for i, (words, norms) in enumerate(zip(tokens, lemmas)):
        for word in set(words):
            postings[word].add(i)
        for word in set(norms):
            lemma_postings[word].add(i)
    df.attrs.update(indexed_texts=texts, tokens=tokens, lemmas=lemmas,
                    postings=dict(postings), lemma_postings=dict(lemma_postings))
    return df


@st.cache_data(show_spinner=False, max_entries=8)
def cached_embeddings(texts, model_id=MODEL_ID):
    """Ключ — тексты, а не параметры отображения. Длинные кейсы режутся по токенам."""
    with model_lock():
        model = get_model(model_id)
        max_tokens = max(8, int(model.max_seq_length) - model.tokenizer.num_special_tokens_to_add(pair=False))
        overlap = min(32, max_tokens // 4)
        chunks, owners = [], []
        for i, text in enumerate(texts):
            ids = model.tokenizer.encode(text, add_special_tokens=False)
            for start in range(0, max(1, len(ids)), max_tokens - overlap):
                chunks.append(model.tokenizer.decode(ids[start:start + max_tokens], skip_special_tokens=True))
                owners.append(i)
                if start + max_tokens >= len(ids):
                    break
        matrix = model.encode(chunks, convert_to_numpy=True, normalize_embeddings=True,
                              batch_size=32, show_progress_bar=False)
        return np.asarray(matrix, dtype=np.float32), np.asarray(owners, dtype=np.int64)


def attach_embeddings(df, model_id=MODEL_ID):
    if not df.empty:
        matrix, owners = cached_embeddings(df.attrs["indexed_texts"], model_id)
        df.attrs.update(phrase_embs=matrix, embedding_owners=owners, model_id=model_id)
    return df


@st.cache_data(show_spinner=False, max_entries=256)
def encode_query(query, model_id=MODEL_ID):
    with model_lock():
        return np.asarray(get_model(model_id).encode(normalize_spaces(query), convert_to_numpy=True,
                          normalize_embeddings=True, show_progress_bar=False), dtype=np.float32)


def semantic_scores(query, df):
    if df.empty or not tokenize(query) or "phrase_embs" not in df.attrs:
        return None
    vector = encode_query(query, df.attrs.get("model_id", MODEL_ID))
    scores = np.clip(df.attrs["phrase_embs"] @ vector.reshape(-1), -1.0, 1.0)
    result = np.full(len(df), -1.0, dtype=np.float32)
    np.maximum.at(result, df.attrs.get("embedding_owners", np.arange(len(df))), scores)
    return result


def _result_from_row(row, score=None, **details):
    result = {key: row[key] for key in ("case_uid", "source_file", "sheet_name", "row_number", "title", "fields", "search_text")}
    result["row_number"] = int(result["row_number"])
    if score is not None:
        result["score"] = float(score)
    result.update(details)
    return result


def deduplicate_results(results, top_k):
    best = {}
    for item in results:
        prev = best.get(item["case_uid"])
        if prev is None or item.get("score", 1) > prev.get("score", 1):
            best[item["case_uid"]] = item
    return sorted(best.values(), key=lambda item: (-item.get("score", 1), item["case_uid"]))[:max(0, top_k)]


def _contains_phrase(words, query_words):
    return any(words[start:start + len(query_words)] == query_words
               for start in range(len(words) - len(query_words) + 1))


def lexical_scores(query, df, allow_substring=False):
    result = np.zeros(len(df), dtype=np.float32)
    kinds = {}
    words = tokenize(query)
    if not words or df.empty:
        return result, kinds
    norms = tuple(map(lemmatize_cached, words))
    attrs = df.attrs
    tokens, lemmas = attrs["tokens"], attrs["lemmas"]
    exact = set.intersection(*(attrs["postings"].get(w, set()) for w in set(words)))
    inflected = set.intersection(*(attrs["lemma_postings"].get(w, set()) for w in set(norms)))
    candidates = exact | inflected
    for i in candidates:
        if _contains_phrase(tokens[i], words):
            result[i], kinds[i] = 3.0, "Фраза целиком"
        elif not (Counter(words) - Counter(tokens[i])):
            result[i], kinds[i] = 2.0, "Все слова"
        elif not (Counter(norms) - Counter(lemmas[i])):
            result[i], kinds[i] = 1.0, "Словоформы"
    if allow_substring:
        texts = attrs["indexed_texts"]
        for i, text in enumerate(texts):
            if result[i] == 0 and preprocess(query) in preprocess(text):
                result[i], kinds[i] = 0.3, "Часть слова (слабый сигнал)"
    return result, kinds


def search_bundle(query, df, top_k=5, threshold=0.5, semantic_weight=0.65,
                  allow_substring=False, use_semantic=True):
    if not 0 <= semantic_weight <= 1 or not -1 <= threshold <= 1:
        raise ValueError("Некорректные параметры поиска.")
    lex, kinds = lexical_scores(query, df, allow_substring)
    sem = semantic_scores(query, df) if use_semantic else None
    hybrid, semantic, lexical = [], [], []
    for i in range(len(df)):
        lscore = float(lex[i])
        sscore = float(sem[i]) if sem is not None else None
        if lscore > 0:
            lexical.append(_result_from_row(df.iloc[i], lscore, match_type=kinds[i]))
        if sscore is not None and sscore >= threshold:
            semantic.append(_result_from_row(df.iloc[i], sscore))
        if not (lscore > 0 or (sscore is not None and sscore >= threshold)):
            continue
        if sem is None:
            combined = lscore / 3
        else:
            combined = semantic_weight * max(0, sscore) + (1 - semantic_weight) * lscore / 3
        if lscore == 3:
            combined += 0.15
        hybrid.append(_result_from_row(df.iloc[i], combined, semantic_score=sscore,
                      lexical_score=lscore, match_type=kinds.get(i, "Семантика")))
    return {"hybrid": deduplicate_results(hybrid, top_k),
            "semantic": deduplicate_results(semantic, top_k),
            "lexical": deduplicate_results(lexical, top_k)}


def semantic_search(query, df, top_k=5, threshold=0.5):
    return search_bundle(query, df, top_k, threshold)["semantic"]


def keyword_search(query, df, top_k=5):
    return search_bundle(query, df, top_k, use_semantic=False)["lexical"]


def hybrid_search(query, df, top_k=5, threshold=0.5, semantic_weight=0.65):
    return search_bundle(query, df, top_k, threshold, semantic_weight)["hybrid"]
